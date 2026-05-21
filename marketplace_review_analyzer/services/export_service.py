from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt

from models import ProductReport
from services.report_generator import category_display, generate_text_report


def export_txt(text: str, output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def export_docx(text: str, output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    for block in text.split("\n"):
        paragraph = document.add_paragraph(block)
        if block.endswith(":") or block.startswith("MAKE UP"):
            for run in paragraph.runs:
                run.bold = True
    document.save(path)
    return path


def export_xlsx(reports: list[ProductReport], output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows: list[dict[str, object]] = []
    for report in reports:
        errors = "; ".join(report.errors)
        base = {
            "Категория": category_display(report.product.category),
            "Продукт": report.product.name,
            "Вариант / оттенок": report.product.variant or "",
            "WB артикул": report.product.wb_article or "",
            "WB ссылка": report.product.wb_url or "",
            "WB рейтинг": report.wb_rating if report.wb_rating is not None else "нет данных",
            "Ozon артикул": report.product.ozon_article or "",
            "Ozon ссылка": report.product.ozon_url or "",
            "Ozon рейтинг": report.ozon_rating if report.ozon_rating is not None else "нет данных",
            "Ошибки обработки": errors,
        }
        if report.problem_groups:
            for group in report.problem_groups:
                rows.append({
                    **base,
                    "Проблема": group.problem,
                    "Количество негативных отзывов": group.count,
                    "Примеры отзывов": "\n---\n".join(group.examples),
                })
        else:
            rows.append({
                **base,
                "Проблема": "Нет новых негативных отзывов",
                "Количество негативных отзывов": 0,
                "Примеры отзывов": "",
            })
    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Отчёт")
        ws = writer.book["Отчёт"]
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 45)
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
    return path


def export_all(reports: list[ProductReport], date_from, date_to, output_dir: str | Path = "outputs") -> dict[str, Path]:
    output_dir = Path(output_dir)
    text = generate_text_report(reports, date_from, date_to)
    return {
        "txt": export_txt(text, output_dir / "report.txt"),
        "docx": export_docx(text, output_dir / "report.docx"),
        "xlsx": export_xlsx(reports, output_dir / "report.xlsx"),
    }
